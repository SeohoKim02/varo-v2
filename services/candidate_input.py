"""Bounded XLSX/CSV adapter for the opt-in economic review, never production.

The original nine sheets remain valid. Four optional policy tables express
period/location/lot facts absent from the canonical contract. Money field names
reuse seller_loss_inputs; no missing cost, reservation or unit is imputed.
"""
from __future__ import annotations

import csv
from dataclasses import dataclass, field
from datetime import date, datetime
import hashlib
import io
import logging
import math
from pathlib import Path
from typing import Any, Mapping
import zipfile

import pandas as pd

from services.candidate_economics import (DemandMarket, EconomicProblem, InventoryLot,
    Money, PriceOption, Route, Scenario, solve_snapshot_economics, validate_problem)
from services.operating_snapshot import from_canonical
from services.operator_intake_candidate import ALLOWED_TABLES, validate_operator_intake
from services.seller_loss_engine import known

class _InputError(ValueError):
    """A deliberately written, user-facing validation message."""


VERSION = "economic-input-adapter-0.1.2"
POLICY_TABLES = {
    "economic_stock_policy": ("location_id", "product_id", "product_state", "lot_id", "quantity_unit",
        "capacity_per_unit", "quantum", "holding_cost_per_unit_day", "disposal_cost_per_unit",
        "salvage_value_per_unit", "currency", "valid_from", "valid_to"),
    "economic_route_policy": ("route_id", "source_id", "target_id", "transit_days", "vehicle_capacity",
        "max_vehicles", "fixed_cost_per_vehicle", "transfer_cost_per_capacity", "allowed_states",
        "departure_periods", "route_approved", "currency", "capacity_unit", "valid_from", "valid_to"),
    "economic_sales_policy": ("location_id", "product_id", "product_state", "quantity_unit", "period",
        "price_level", "unit_price", "demand_qty", "scenario_id", "probability", "demand_basis",
        "currency", "valid_from", "valid_to"),
    "economic_location_policy": ("location_id", "capacity_unit", "storage_capacity",
        "handling_cost_per_capacity", "processing_capacity", "currency", "valid_from", "valid_to"),
}
OPTIONAL_TABLES = {
    "economic_flow_policy": ("event_id", "lot_id", "location_id", "period", "quantity", "kind", "quantity_unit"),
    "economic_commitment_policy": ("commitment_id", "lot_id", "route_id", "departure_period", "quantity", "vehicle_count", "cost_already_paid", "quantity_unit"),
}
TABLE_LABELS = {"inventory_snapshot": "재고 현황", "economic_stock_policy": "상품별 비용",
    "economic_route_policy": "이동 경로", "economic_sales_policy": "판매가격과 예상 수요",
    "economic_location_policy": "점포 용량", "metadata": "입력 설정"}
FIELD_HELP = {
    "holding_cost_per_unit_day": "수량 1단위를 하루 보관하는 비용 (통화/수량/일)",
    "disposal_cost_per_unit": "수량 1단위를 폐기할 때 추가 지출하는 비용",
    "salvage_value_per_unit": "계획 마지막 날 남은 수량 1단위의 회수 가치",
    "capacity_per_unit": "수량 1단위가 차지하는 공통 용량 (예: kg/개)",
    "quantum": "나눌 수 있는 최소 수량 (개별 상품은 1)",
    "approved_transfer_qty": "실제 확인된 최대 이동 승인량. 현재 재고와 다릅니다",
    "price_level": "normal은 정상판매, 그 외는 명시한 할인 정책 이름",
    "demand_basis": "FORECAST는 예상 수요, SCENARIO는 시험 가정. 관측 판매를 미래 수요로 쓰지 않음",
    "departure_periods": "출발 가능한 날짜 순서. 첫날 0, 다음날 1. 여러 값은 쉼표로 구분",
    "valid_from": "가격·비용·용량 근거가 적용되는 첫 날짜 (YYYY-MM-DD)",
    "valid_to": "계획 종료일까지 유효해야 하는 근거의 마지막 날짜",
    "reserved_qty": "다른 주문에 예약된 수량. 예약 없음도 명시적으로 0 입력",
    "confirmed_inbound_qty": "확정 입고량. 0이 아니면 도착 일정 별도 필요",
    "confirmed_outbound_qty": "확정 출고량. 0이 아니면 출고 일정 별도 필요",
}


@dataclass
class InputBundle:
    tables: dict[str, pd.DataFrame] = field(default_factory=dict)
    hashes: dict[str, str] = field(default_factory=dict)
    issues: list[dict[str, Any]] = field(default_factory=list)


@dataclass
class PreparedInput:
    issues: list[dict[str, Any]]
    data_class: str
    snapshot: Any = None
    problem: EconomicProblem | None = None
    file_hashes: dict[str, str] = field(default_factory=dict)

    @property
    def ready(self):
        return self.problem is not None and not self.issues


def _issue(table, row, column, message):
    return {"table": table, "row": row, "column": column, "message_ko": message}


def _blank(value):
    return value is None or str(value).strip() == "" or (isinstance(value, float) and math.isnan(value))


def _text(value):
    return "" if _blank(value) else str(value).strip()


def _date(value):
    if isinstance(value, datetime): return value.date()
    if isinstance(value, date): return value
    return date.fromisoformat(str(value).strip())


def _boolean(value):
    if value is True or _text(value).lower() in {"true", "1", "예", "확인"}: return True
    if value is False or _text(value).lower() in {"false", "0", "아니오"}: return False
    return None


def read_uploads(files: Mapping[str, bytes], *, max_bytes=16 * 1024 * 1024) -> InputBundle:
    """Read original bytes only. XLSX formulas are blocked, even with a cache:
    OOXML does not prove cache freshness, so an explicit values-only export is required.
    Never execute macros, external links or spreadsheet formulas.
    """
    out = InputBundle()
    for filename, raw in files.items():
        if not raw or len(raw) > max_bytes:
            out.issues.append(_issue(filename, None, None, "파일이 비었거나 16MB를 넘습니다. 기간을 나누어 다시 저장해 주세요.")); continue
        digest = hashlib.sha256(raw).hexdigest()
        out.hashes[filename] = digest
        suffix = Path(filename).suffix.lower()
        tables = {}
        try:
            if suffix == ".csv":
                table = Path(filename).stem
                text = raw.decode("utf-8-sig")
                rows = list(csv.reader(io.StringIO(text)))
                tables[table] = rows
            elif suffix == ".xlsx":
                with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                    if sum(z.file_size for z in archive.infolist()) > 64 * 1024 * 1024:
                        raise _InputError("압축을 푼 엑셀 자료가 너무 큽니다. 기간을 나눠 주세요.")
                    if any("externalLinks/" in z.filename or "vbaProject" in z.filename for z in archive.infolist()):
                        raise _InputError("외부 연결이나 매크로가 있습니다. 값만 복사한 새 XLSX 파일을 사용해 주세요.")
                from openpyxl import load_workbook
                book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
                try:
                    for sheet in book:
                        if sheet.title in {"작성 안내", "필드 안내"}: continue
                        if (sheet.max_row or 0) > 100001 or (sheet.max_column or 0) > 200:
                            raise _InputError("표가 너무 큽니다. 10만 행·200열 이내로 기간을 나누어 주세요.")
                        values = []
                        for row_number, row in enumerate(sheet.iter_rows(), 1):
                            if row_number > 100001 or len(row) > 200:
                                raise _InputError("표가 너무 큽니다. 10만 행·200열 이내로 기간을 나누어 주세요.")
                            cells = []
                            for c in row:
                                if c.data_type == "f":
                                    out.issues.append(_issue(sheet.title, c.row, c.column_letter,
                                        "수식 결과의 최신 여부를 확인할 수 없습니다. 엑셀에서 다시 계산한 뒤 값만 붙여넣어 저장해 주세요."))
                                    cells.append(None)
                                else: cells.append(c.value)
                            values.append(cells)
                        tables[sheet.title] = values
                finally: book.close()
            else: raise _InputError("XLSX 또는 UTF-8 CSV만 사용할 수 있습니다.")
            for table, rows in tables.items():
                if not rows or all(all(_blank(x) for x in row) for row in rows): continue
                if table not in ALLOWED_TABLES and table not in POLICY_TABLES and table not in OPTIONAL_TABLES:
                    out.issues.append(_issue(table, 1, None, "지원하지 않는 표입니다. 기존 입력 양식의 표 이름을 확인해 주세요.")); continue
                if table in out.tables:
                    out.issues.append(_issue(table, None, None, "같은 표가 여러 파일에 있습니다. 자동으로 합치지 않았으니 하나의 표로 정리해 주세요.")); continue
                header = [_text(x) for x in rows[0]]
                while header and not header[-1]: header.pop()
                if not header or any(not x for x in header) or len(header) != len(set(header)):
                    out.issues.append(_issue(table, 1, None, "열 이름이 비어 있거나 중복됩니다. 첫 행을 확인해 주세요.")); continue
                records = []
                for n, row in enumerate(rows[1:], 2):
                    if not any(not _blank(x) for x in row): continue
                    if len(row) > len(header) and any(not _blank(x) for x in row[len(header):]):
                        out.issues.append(_issue(table, n, None, "제목 없는 열에 값이 있습니다. 열 이름을 추가하거나 값을 정리해 주세요.")); continue
                    rec = dict(zip(header, [*row, *([None] * len(header))][:len(header)]))
                    rec["_input_row"] = n
                    rec["_input_file"] = filename
                    records.append(rec)
                out.tables[table] = pd.DataFrame(records, columns=[*header, "_input_row", "_input_file"])
        except (ValueError, KeyError, zipfile.BadZipFile, UnicodeDecodeError, OSError) as exc:
            logging.getLogger(__name__).warning("Candidate input file rejected: %s", filename, exc_info=True)
            out.issues.append(_issue(filename, None, None, str(exc) if isinstance(exc, _InputError) else "파일을 읽지 못했습니다. XLSX 또는 UTF-8 CSV로 다시 저장해 주세요."))
    return out


def prepare_input(bundle: InputBundle, metadata: Mapping[str, Any]) -> PreparedInput:
    """Build one network/day problem; failures preserve all rows for correction."""
    issues = list(bundle.issues)
    kind = metadata.get("data_class", "REAL")
    result = PreparedInput(issues, kind, file_hashes=dict(bundle.hashes))
    if kind not in {"REAL", "TEST", "SCENARIO"}:
        issues.append(_issue("metadata", None, None, "실제 자료·시험 입력·가상 시나리오 중 하나를 선택해 주세요."))
    if metadata.get("owner_authorized") is not True:
        issues.append(_issue("metadata", None, None, "자료 제공자의 연구 검토 이용 허락을 확인해 주세요."))
    dataset, network = _text(metadata.get("dataset_id")), _text(metadata.get("network_id"))
    try:
        as_of = _date(metadata.get("as_of_date")); horizon = int(metadata.get("horizon_days"))
        if not 1 <= horizon <= 31: raise ValueError()
    except (ValueError, TypeError):
        issues.append(_issue("metadata", None, None, "기준일과 1~31일의 계획기간을 입력해 주세요.")); return result
    if not dataset or not network:
        issues.append(_issue("metadata", None, None, "자료 묶음과 사업장 식별 이름이 필요합니다."))
    currency, capunit = _text(metadata.get("currency")), _text(metadata.get("capacity_unit"))
    if not currency or not capunit: issues.append(_issue("metadata", None, None, "통화와 공통 용량 단위를 확인해 주세요."))
    frames = {k: v.copy() for k, v in bundle.tables.items()}
    inv = frames.get("inventory_snapshot", pd.DataFrame())
    if inv.empty:
        issues.append(_issue("inventory_snapshot", None, None, "재고 기록이 없습니다. 재고 현황 표에 실제 행을 입력해 주세요.")); return result
    records = inv.astype(object).where(pd.notna(inv), None).to_dict("records")
    for r in records:
        for name in ("analysis_eligible", "lot_required"):
            if name in r: r[name] = _boolean(r[name])
        for name in ("snapshot_date", "expiry_date", "transfer_valid_until"):
            if not _blank(r.get(name)):
                try: r[name] = _date(r[name]).isoformat()
                except ValueError: issues.append(_issue("inventory_snapshot", r["_input_row"], name, "날짜 형식이 맞지 않습니다. YYYY-MM-DD로 입력해 주세요."))
        for raw_id, key in (("location_id", "location_key"), ("product_id", "product_key")):
            ns = f"{network}:{raw_id}"
            if _blank(r.get(raw_id)): issues.append(_issue("inventory_snapshot", r["_input_row"], raw_id, "점포와 상품의 식별 코드가 필요합니다."))
            if _blank(r.get(raw_id + "_namespace")): r[raw_id + "_namespace"] = ns
            if _blank(r.get(key)): r[key] = f"{ns}:{_text(r.get(raw_id))}"
            if r[raw_id + "_namespace"] != ns or r[key] != f"{ns}:{_text(r.get(raw_id))}":
                issues.append(_issue("inventory_snapshot", r["_input_row"], raw_id, "식별 코드의 사업장 범위가 다릅니다. 서로 다른 사업장 자료를 섞지 마세요."))
        # Physical upload lineage is appended, never used to replace a conflicting supplied source.
        for name, value in {"source_file": r["_input_file"], "source_sheet_or_table": "inventory_snapshot",
                            "source_row_id": str(r["_input_row"]), "transform_version": VERSION,
                            "source_dataset": dataset}.items():
            if _blank(r.get(name)): r[name] = value
    if kind == "REAL":
        checked = validate_operator_intake({"inventory_snapshot": pd.DataFrame(records)},
            {**metadata, "allowed_use": "RESEARCH_SHADOW", "publication_scope": "NOT_ALLOWED", "source_hashes": bundle.hashes})
        for i in checked.issues:
            if i["severity"] == "ERROR":
                issues.append(_issue(i["table"], i["row"] + 2 if i["row"] is not None else None,
                    ", ".join(i["fields"]), i["message_ko"]))
    snapshot = from_canonical(records, dataset=dataset, network=network, as_of=as_of.isoformat(), horizon_days=horizon)
    result.snapshot = snapshot
    for p, row in zip(snapshot.positions, records):
        if p.issues:
            issues.append(_issue("inventory_snapshot", row["_input_row"], None,
                "점포·상품·배치·단위·재고·이동 승인 근거를 확인해 주세요."))

    def rows(table):
        frame = frames.get(table, pd.DataFrame())
        if frame.empty:
            issues.append(_issue(table, None, None, f"{TABLE_LABELS[table]} 표가 필요합니다. 미확인 비용을 0으로 계산하지 않습니다."))
            return []
        values = frame.astype(object).where(pd.notna(frame), None).to_dict("records")
        for row in values:
            for col in POLICY_TABLES[table]:
                if _blank(row.get(col)):
                    issues.append(_issue(table, row["_input_row"], col, FIELD_HELP.get(col, "필수 값을 원본 근거와 함께 입력해 주세요.")))
            try:
                end = as_of.fromordinal(as_of.toordinal() + horizon - 1)
                if not _date(row.get("valid_from")) <= as_of <= end <= _date(row.get("valid_to")): raise ValueError()
            except (TypeError, ValueError): issues.append(_issue(table, row["_input_row"], "valid_from / valid_to", "가격·비용·용량의 적용기간이 계획기간을 포함해야 합니다."))
            if _text(row.get("currency")) != currency: issues.append(_issue(table, row["_input_row"], "currency", "통화가 입력 설정과 다릅니다. 자동 환산하지 않습니다."))
        return values

    stockrows, routerows, salesrows, locrows = (rows(t) for t in POLICY_TABLES)
    if issues: return result
    def num(row, name, *, integer=False):
        value = row.get(name)
        try:
            n = float(value)
            if isinstance(value, bool) or not math.isfinite(n) or n < 0 or (integer and n != int(n)): raise ValueError()
            return int(n) if integer else n
        except (ValueError, TypeError): raise _InputError("수량·금액·기간에 0 이상의 올바른 숫자가 필요합니다. 수량 구분과 단위를 함께 확인해 주세요.")
    def identity(row): return tuple(_text(row.get(k)) for k in ("location_id", "product_id", "product_state", "lot_id"))
    def money(row, name, unit):
        return Money(known(num(row, name), "USER_INPUT" if kind == "REAL" else "SCENARIO_INPUT",
            f"{row['_input_file']}:{row['_input_row']}:{name}", unit=unit, currency=currency, dataset=dataset), "KNOWN")
    def lid(raw): return f"{network}:location_id:{_text(raw)}"
    def pid(raw): return f"{network}:product_id:{_text(raw)}"
    try:
        by_stock = {identity(r): r for r in stockrows}
        if len(by_stock) != len(stockrows): raise _InputError("상품별 비용 표의 식별 코드가 중복됩니다.")
        committed_lots = set(_text(r.get("lot_id")) for r in frames.get("economic_commitment_policy", pd.DataFrame()).to_dict("records"))
        lots = []
        for r, position in zip(records, snapshot.positions):
            if position.on_hand == 0 and not position.inbound_confirmed and position.lot not in committed_lots: continue
            p = by_stock.get(identity(r))
            if p is None: raise _InputError("일부 재고의 상품별 비용 정보가 없습니다. 모든 상품·배치의 비용을 확인해 주세요.")
            if p["quantity_unit"] != position.unit: raise _InputError("재고와 비용의 수량 단위가 다릅니다.")
            if not position.expiry_date: raise _InputError("배치별 유통기한이 필요합니다.")
            lots.append(InventoryLot(position.lot, position.product, position.location, position.state, position.unit,
                position.on_hand, (_date(position.expiry_date)-as_of).days, num(p, "capacity_per_unit"),
                money(p, "holding_cost_per_unit_day", position.unit), money(p, "disposal_cost_per_unit", position.unit),
                money(p, "salvage_value_per_unit", position.unit), num(p, "quantum"), position.approved_transfer))
        routes = []
        for r in routerows:
            if _boolean(r["route_approved"]) is not True: raise _InputError("경로 사용 승인 여부를 확인해 주세요.")
            if r["capacity_unit"] != capunit: raise _InputError("경로 용량의 단위가 공통 용량 단위와 다릅니다.")
            try:
                dep = tuple(int(x.strip()) for x in _text(r["departure_periods"]).split(","))
            except (TypeError, ValueError) as exc:
                logging.getLogger(__name__).warning("Invalid route departure periods", exc_info=True)
                raise _InputError("출발 가능한 날짜 순서를 정수로 입력해 주세요. 첫날은 0, 다음날은 1이며 여러 날짜는 쉼표로 구분합니다.") from exc
            routes.append(Route(_text(r["route_id"]), lid(r["source_id"]), lid(r["target_id"]), num(r,"transit_days",integer=True),
                num(r,"vehicle_capacity"), num(r,"max_vehicles",integer=True), money(r,"fixed_cost_per_vehicle",None),
                money(r,"transfer_cost_per_capacity",capunit), tuple(x.strip() for x in str(r["allowed_states"]).split(",")), dep))
        groups, probabilities, signals = {}, {}, {}
        for r in salesrows:
            if r["demand_basis"] not in {"FORECAST", "SCENARIO"}: raise _InputError("미래 수요는 예측 또는 시나리오로 명시해야 합니다. 관측 판매를 미래의 정답으로 사용하지 않습니다.")
            if kind == "REAL" and r["demand_basis"] != "FORECAST": raise _InputError("가정 수요를 포함하면 자료 구분을 시험 입력 또는 가상 시나리오로 변경해 주세요.")
            sid, prob = _text(r["scenario_id"]), num(r,"probability")
            if sid in probabilities and probabilities[sid] != prob: raise _InputError("같은 수요 시나리오의 발생 비율이 다릅니다.")
            probabilities[sid] = prob
            # Optional columns extend the existing sales policy; no new mandatory sheet.
            signal_period, signal = r.get("observation_period"), r.get("observation_signal")
            if not _blank(signal_period) or not _blank(signal):
                if _blank(signal_period) or _blank(signal):
                    raise _InputError("수요 상황의 관측 시점과 내용을 함께 입력해 주세요.")
                t = num(r, "observation_period", integer=True)
                if t >= horizon: raise _InputError("수요 상황의 관측 시점은 계획기간 안이어야 합니다.")
                scenario_signals = signals.setdefault(sid, {})
                if t in scenario_signals and scenario_signals[t] != _text(signal):
                    raise _InputError("같은 수요 시나리오의 관측 상황이 서로 다릅니다.")
                scenario_signals[t] = _text(signal)
            key = (lid(r["location_id"]), pid(r["product_id"]), _text(r["product_state"]), _text(r["quantity_unit"]), num(r,"period",integer=True))
            options = groups.setdefault(key, {})
            level = _text(r["price_level"])
            option = options.setdefault(level, {"price": money(r,"unit_price",key[3]), "demand": {}})
            if option["price"].field.value != num(r,"unit_price") or sid in option["demand"]: raise _InputError("판매가격 또는 동일 시나리오의 수요 기록이 중복·충돌합니다.")
            option["demand"][sid] = num(r,"demand_qty")
        markets = tuple(DemandMarket(f"market_{n}", *key, tuple(PriceOption(level, v["price"], v["demand"]) for level,v in options.items()))
            for n,(key,options) in enumerate(groups.items()))
        storage, handling, processing = {}, {}, {}
        for r in locrows:
            key = lid(r["location_id"])
            if key in storage or r["capacity_unit"] != capunit: raise _InputError("점포 용량 기록이 중복되거나 단위가 다릅니다.")
            storage[key] = num(r,"storage_capacity")
            handling[key] = money(r,"handling_cost_per_capacity",capunit)
            processing.update({(key,t): num(r,"processing_capacity") for t in range(horizon)})
        from dataclasses import fields
        extras = {}
        available = {f.name for f in fields(EconomicProblem)}
        if "handling_per_capacity" in available:
            extras.update(handling_per_capacity=handling, processing_capacity=processing)
        elif any(x.field.value for x in handling.values()): raise _InputError("취급비 모형 연결을 완료한 뒤 계산해 주세요.")
        from services.candidate_economics import ConfirmedFlow, CommittedTransfer
        lots_by_id = {l.lot_id: l for l in lots}
        flow_events, commitments = [], []
        for table, required in OPTIONAL_TABLES.items():
            frame = frames.get(table, pd.DataFrame())
            for row in frame.astype(object).where(pd.notna(frame), None).to_dict("records"):
                if any(_blank(row.get(col)) for col in required):
                    raise _InputError("확정 입출고·운송 중 재고의 일정과 수량을 빠짐없이 입력해 주세요.")
                lot = lots_by_id.get(_text(row["lot_id"]))
                if lot is None or lot.unit != row["quantity_unit"]:
                    raise _InputError("확정 입출고의 배치 또는 수량 단위가 재고와 다릅니다.")
                if table == "economic_flow_policy":
                    flow_events.append(ConfirmedFlow(_text(row["event_id"]), lot.lot_id, lid(row["location_id"]),
                        num(row,"period",integer=True), num(row,"quantity"), _text(row["kind"])))
                else:
                    try:
                        departed = float(row["departure_period"])
                    except (TypeError, ValueError) as exc:
                        logging.getLogger(__name__).warning("Invalid committed departure period", exc_info=True)
                        raise _InputError("확정 이동의 출발 시점을 정수로 입력해 주세요. 계획 시작 전 출발은 음수로 표시합니다.") from exc
                    paid = _boolean(row["cost_already_paid"])
                    if not math.isfinite(departed) or departed != int(departed) or paid is None:
                        raise _InputError("확정 이동의 출발 시점과 운송비 지출 여부를 확인해 주세요.")
                    commitments.append(CommittedTransfer(_text(row["commitment_id"]), lot.lot_id, _text(row["route_id"]),
                        int(departed), num(row,"quantity"), num(row,"vehicle_count",integer=True), paid))
        extras.update(confirmed_flows=tuple(flow_events), committed_transfers=tuple(commitments))
        problem = EconomicProblem(f"{dataset}:{as_of}", horizon, currency, tuple(lots), tuple(routes), markets,
            tuple(Scenario(k,v, observations=signals.get(k, {})) for k,v in probabilities.items()), capunit, storage,
            evidence_class="REAL_MODEL" if kind == "REAL" else "SCENARIO",
            lineage={"adapter": VERSION, "input_data_class": kind, **bundle.hashes}, **extras)
        from services.candidate_economics_extended import validate_extension
        for issue in validate_extension(problem):
            issues.append(_issue("metadata", None, None, "운영 모형에 필요한 기간·가격·수요·단위·용량을 확인해 주세요."))
        if not issues:
            from dataclasses import replace
            result.snapshot = replace(snapshot, context={"approved_route_ids": [r.route_id for r in routes], "input_data_class": kind})
            from services.candidate_economics import bind_snapshot_economics
            bound, binding_issues = bind_snapshot_economics(result.snapshot, problem)
            if binding_issues:
                logging.getLogger(__name__).warning("Candidate snapshot binding rejected: %r", binding_issues)
                issues.append(_issue("metadata", None, None, "초기 재고·예약량·확정 입출고 합계와 일정이 일치하지 않습니다. 수량과 근거를 확인해 주세요."))
            else:
                result.problem = bound
    except (ValueError, TypeError, KeyError) as exc:
        logging.getLogger(__name__).warning("Candidate input model rejected", exc_info=True)
        issues.append(_issue("metadata", None, None, str(exc) if isinstance(exc,_InputError) else "숫자·식별 코드·필수 열을 확인해 주세요."))
    return result


def run_prepared(prepared: PreparedInput, *, include_comparison: bool = True) -> dict[str, Any]:
    if not prepared.ready: return {"status": "INPUT_REQUIRED", "input_issues": prepared.issues, "production_approved": False}
    result = solve_snapshot_economics(prepared.snapshot, prepared.problem)
    result["input_data_class"] = prepared.data_class
    result["input_file_hashes"] = prepared.file_hashes
    result["actual_savings"] = None
    result["production_approved"] = False
    if include_comparison and result.get("status") in {"OPTIMAL", "FEASIBLE_LIMIT"}:
        from services.candidate_economics_policies import fair_baselines
        comparison = fair_baselines(prepared.problem)
        result["alternative_comparison"] = comparison
    return result
