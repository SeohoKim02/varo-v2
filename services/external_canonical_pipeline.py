"""Canonical exports of the external generalisation datasets (schema 1.2.0) and the validation matrix.

Run through the shared entry point (raw is read-only; outputs are processed/canonical_*.parquet and results/*):
    python -m services.canonical_data_pipeline --data-root C:/VARO_V2_REAL_DATA --datasets m5 favorita freshretailnet kamp
Rebuild only _COLLECTION_STATUS/EXTERNAL_VALIDATION_MATRIX.csv from existing outputs:
    python -m services.external_canonical_pipeline --data-root C:/VARO_V2_REAL_DATA --matrix-only
Large sources are streamed: M5 500 series (x1,941 days) per chunk, Favorita 64 MB CSV blocks,
FreshRetailNet 250,000-row parquet batches. Key uniqueness of those tables is checked with one
64-bit hash per row (Exporter.bulk_key_tables).
"""
from __future__ import annotations

import argparse
import json
import math
from collections import Counter
from pathlib import Path

import numpy as np
import pandas as pd

from services.canonical_data_pipeline import ALGORITHMS
from services.canonical_schema import EXTERNAL_VERSION, NUMERIC
from services.external_data_adapters import (EXTERNAL_UNIT_EVIDENCE, FRN_CATEGORY_LEVELS, FRN_HOURS, FRN_STOCKOUT_WINDOW, FRN_WEATHER, KAMP_GRADES,
    KAMP_KEYS, KAMP_TOTAL, M5_META, M5_STATES, NAMESPACES, adapt_favorita_holidays, adapt_favorita_masters, adapt_favorita_oil,
    adapt_favorita_train, adapt_favorita_transactions, adapt_frn_covariates, adapt_frn_masters, adapt_frn_sales, adapt_kamp_shipments,
    adapt_m5_calendar, adapt_m5_masters, adapt_m5_sales)
from services.real_data_adapters import DATA_ROOT, DATASETS, csv_chunks, csv_encoding

IDENTITY_DERIVED = {
    "product_key": {"source_columns": ["product_id_namespace", "product_id"], "formula": "namespace + ':' + product_id (canonicalize)", "unit": None, "assumptions": "Cross-dataset identity; raw id unchanged"},
    "location_key": {"source_columns": ["location_id_namespace", "location_id"], "formula": "namespace + ':' + location_id (canonicalize)", "unit": None, "assumptions": "Cross-dataset identity; raw id unchanged"},
}
FLAG_DERIVED = {name: {"source_columns": ["(all mapped columns)"], "formula": "canonical schema validation", "unit": None, "assumptions": "Flags do not repair observations"}
                for name in ("validation_flags", "quality_flags", "quality_status", "analysis_eligible")}


def prepare(e, dataset):
    e.transform_version = EXTERNAL_VERSION
    e.unit_evidence = EXTERNAL_UNIT_EVIDENCE[dataset]


def rel(e, path):
    return Path(path).relative_to(e.root).as_posix()


def read_all(path, **kwargs):
    return pd.concat(list(csv_chunks(Path(path), **kwargs)))


def derived(namespace_fields=(), **extra):
    out = {**FLAG_DERIVED, **extra}
    for field, namespace in namespace_fields:
        out[field] = {"source_columns": [], "formula": repr(namespace), "unit": None, "assumptions": "Source dataset and field namespace"}
        key = "product_key" if field.startswith("product") else "location_key"
        out[key] = IDENTITY_DERIVED[key]
    return out


def exact_sum(values):
    values = pd.to_numeric(pd.Series(values), errors="coerce").dropna()
    return math.fsum(values.to_numpy(dtype=float))


class StreamConservation:
    """Streamed conservation of one source quantity: exact (fsum) chunk sums, non-null counts and record/cell counts."""

    def __init__(self, source_file, table, field, source_column, method):
        self.entry = {"source_file": source_file, "table": table, "field": field, "source_column": source_column, "method": method}
        self.source_sums, self.canonical_sums = [], []
        self.counts = Counter()

    def add(self, source_values, canonical_values, source_records, canonical_rows, source_cells=None):
        source = pd.to_numeric(pd.Series(np.asarray(source_values, dtype=object)), errors="coerce")
        canonical = pd.to_numeric(pd.Series(canonical_values).reset_index(drop=True), errors="coerce")
        self.source_sums.append(exact_sum(source))
        self.canonical_sums.append(exact_sum(canonical))
        self.counts.update(rows_source=source_records, rows_canonical=canonical_rows, source_cells=len(source) if source_cells is None else source_cells,
                           source_nonnull=int(source.notna().sum()), canonical_nonnull=int(canonical.notna().sum()),
                           source_zero=int(source.eq(0).sum()), canonical_zero=int(canonical.eq(0).sum()))

    def result(self):
        a, b = math.fsum(self.source_sums), math.fsum(self.canonical_sums)
        c = self.counts
        # Sums, non-null counts and zero counts must all match: NULL and zero stay distinct through the transform.
        passed = (math.isclose(a, b, rel_tol=1e-12, abs_tol=1e-7) and c["source_nonnull"] == c["canonical_nonnull"]
                  and c["source_zero"] == c["canonical_zero"] and c["source_cells"] == c["rows_canonical"])
        entry = {**self.entry, "rows_source": c["rows_source"], "rows_canonical": c["rows_canonical"], "source_cells": c["source_cells"],
                 "source_records_represented": c["rows_source"], "source_sum": a, "canonical_sum": b,
                 "source_nonnull": c["source_nonnull"], "canonical_nonnull": c["canonical_nonnull"],
                 "canonical_nonnull_records": c["canonical_nonnull"], "source_zero": c["source_zero"], "canonical_zero": c["canonical_zero"],
                 "passed": bool(passed)}
        if not passed:
            raise AssertionError(f"Conservation failure: {entry}")
        return entry


def arrow_csv_chunks(path, block_size=64 << 20):
    """Stream a large UTF-8 CSV as all-string pandas chunks; only an empty field is NULL (same as csv_chunks).

    Each chunk's index is the 0-based data-record position across the whole file.
    """
    import pyarrow as pa
    import pyarrow.csv as pcsv
    if csv_encoding(Path(path)) != "utf-8-sig":
        raise ValueError(f"Arrow streaming requires UTF-8: {path}")
    with open(path, encoding="utf-8-sig") as handle:
        header = handle.readline().strip().split(",")
    convert = pcsv.ConvertOptions(column_types={c: pa.string() for c in header}, null_values=[""], strings_can_be_null=True,
                                  quoted_strings_can_be_null=True)
    reader = pcsv.open_csv(path, read_options=pcsv.ReadOptions(block_size=block_size), convert_options=convert)
    position = 0
    for batch in reader:
        frame = batch.to_pandas()
        frame.index = pd.RangeIndex(position, position + len(frame))
        position += len(frame)
        yield frame


# ---------------------------------------------------------------- M5

def m5_validation_release_evidence(evaluation, validation, chunksize=2000):
    """sales_train_validation.csv must be a strict prefix release of sales_train_evaluation.csv (not converted twice)."""
    rows = mismatched_cells = compared = 0
    same_order = True
    days = None
    for ev, va in zip(csv_chunks(evaluation, chunksize=chunksize), csv_chunks(validation, chunksize=chunksize)):
        days = days or [c for c in va.columns if c.startswith("d_")]
        same_order &= bool((ev[list(M5_META)].to_numpy() == va[list(M5_META)].to_numpy()).all())
        diff = ev[days].to_numpy() != va[days].to_numpy()
        mismatched_cells += int(diff.sum())
        compared += diff.size
        rows += len(va)
    return {"validation_rows": rows, "validation_days": [days[0], days[-1]], "same_series_order": same_order, "cells_compared": compared,
            "cells_different": mismatched_cells, "strict_prefix": same_order and mismatched_cells == 0,
            "policy": "Validation release converted zero times: it repeats d_1..d_1913 of the evaluation release cell for cell; converting both would double-count sales."}


def run_m5(e):
    prepare(e, "m5")
    folder = e.folder / "raw/extracted/m5-forecasting-accuracy"
    paths = {n: folder / f"{n}.csv" for n in ("calendar", "sales_train_evaluation", "sales_train_validation", "sell_prices", "sample_submission")}
    calendar = read_all(paths["calendar"])
    if not (calendar["d"].is_unique and calendar["date"].is_unique):
        raise ValueError("calendar.csv d/date are not one-to-one")
    prices = read_all(paths["sell_prices"], chunksize=2_000_000)
    if prices.duplicated(["store_id", "item_id", "wm_yr_wk"]).any():
        raise ValueError("sell_prices.csv repeats a (store, item, week) price")
    prices["_pair"] = prices["store_id"] + "|" + prices["item_id"]
    source = rel(e, paths["sales_train_evaluation"])
    e.bulk_key_tables.add("demand_series")
    sales = StreamConservation(source, "demand_series", "sales_qty", "d_1..d_1941 (every cell unpivoted)",
                               "fsum of every wide cell vs every canonical row; cells == rows; non-null cells == non-null rows")
    price = Counter()
    price_sums = {"canonical": [], "expected": []}
    first_rows, flags = [], Counter()
    day_week = calendar.set_index("d")["wm_yr_wk"]
    days = horizon = None
    for wide in csv_chunks(paths["sales_train_evaluation"], chunksize=500):
        if days is None:
            days = [c for c in wide.columns if c.startswith("d_")]
            if days != [f"d_{i}" for i in range(1, len(days) + 1)]:
                raise ValueError("d_* columns are not contiguous")
            horizon = day_week.loc[days].value_counts()
        chunk_prices = prices[prices["_pair"].isin(set(wide["store_id"] + "|" + wide["item_id"]))]
        table, frame, long = adapt_m5_sales(wide, source, calendar, chunk_prices)
        e.write(table, frame, masters=False)
        sales.add(wide[days].to_numpy().ravel(), frame["sales_qty"], len(wide), len(frame))
        # Every in-horizon price record lands on exactly the sales days of its week; nothing else is priced.
        in_horizon = chunk_prices[chunk_prices["wm_yr_wk"].isin(horizon.index)]
        repeat = in_horizon["wm_yr_wk"].map(horizon).astype(int).to_numpy()
        price.update(records_in_horizon=len(in_horizon), records_beyond_horizon=len(chunk_prices) - len(in_horizon),
                     expected_priced_rows=int(repeat.sum()), priced_rows=int(frame["price"].notna().sum()),
                     records_joined=int(long.loc[long["sell_price"].notna(), ["store_id", "item_id", "wm_yr_wk"]].drop_duplicates().shape[0]))
        price_sums["expected"].append(math.fsum(np.repeat(pd.to_numeric(in_horizon["sell_price"]).to_numpy(), repeat)))
        price_sums["canonical"].append(exact_sum(frame["price"]))
        for value, n in frame["validation_flags"].value_counts().items():
            flags.update({f: n for f in str(value).split("|") if f})
        first_rows.append(wide[list(M5_META)].assign(_record=(wide.index + 1).astype(str)))
    conservation = sales.result()
    e.conservation.append(conservation)
    price_ok = (price["priced_rows"] == price["expected_priced_rows"] and price["records_joined"] == price["records_in_horizon"]
                and math.isclose(math.fsum(price_sums["canonical"]), math.fsum(price_sums["expected"]), rel_tol=1e-12))
    e.conservation.append({"source_file": rel(e, paths["sell_prices"]), "table": "demand_series", "field": "price", "source_column": "sell_price",
                           "method": "each price record within d_1..d_1941 joins to exactly the sales days of its week (7, or 2 for the last partial week)",
                           "source_records": len(prices), **dict(price), "source_price_x_days_sum": math.fsum(price_sums["expected"]),
                           "canonical_price_sum": math.fsum(price_sums["canonical"]), "passed": bool(price_ok)})
    if not price_ok:
        raise AssertionError(f"M5 price join conservation failure: {dict(price)}")
    first = pd.concat(first_rows)
    for child, parent in [("item_id", "dept_id"), ("dept_id", "cat_id"), ("store_id", "state_id")]:
        if first.groupby(child)[parent].nunique().gt(1).any():
            raise ValueError(f"M5 hierarchy is not a function: {child} -> {parent}")
    products, locations = adapt_m5_masters(first, source)
    e.write("product_master", products, masters=False)
    e.write("location_master", locations, masters=False)
    events, snap = adapt_m5_calendar(calendar, rel(e, paths["calendar"]))
    e.write("calendar_event", events, masters=False)
    e.write("covariate_series", snap, masters=False)
    snap_check = StreamConservation(rel(e, paths["calendar"]), "covariate_series", "covariate_value", "snap_CA|snap_TX|snap_WI", "every snap cell -> one state row")
    snap_check.add(calendar[[f"snap_{s}" for s in M5_STATES]].to_numpy().ravel(order="F"), snap["covariate_value"], len(calendar), len(snap))
    e.conservation.append(snap_check.result())
    expected_events = int(calendar["event_name_1"].notna().sum() + calendar["event_name_2"].notna().sum())
    e.conservation.append({"source_file": rel(e, paths["calendar"]), "table": "calendar_event", "field": "event_name", "source_column": "event_name_1|event_name_2",
                           "source_nonnull": expected_events, "rows_canonical": len(events), "passed": expected_events == len(events)})
    if e.units["demand_series"] == {"item"} and set(e.unit_status["demand_series"]) == {"SOURCE_METADATA"}:
        # Verified pseudo-field for the coverage gate: one documented unit on every sales row.
        e.present.add("consistent_unit")
    ns = NAMESPACES["m5"]
    e.record_mapping("demand_series", {"location_id": "store_id", "product_id": "item_id", "category": "cat_id", "sales_qty": "d_1..d_1941", "raw_date": "d_* header"},
                     source, {"sales_qty": "d_1..d_1941"}, "daily", "store_product", "item",
                     derived([("product_id_namespace", ns["product"]), ("location_id_namespace", ns["store"])],
                             date={"source_columns": ["d_* header", "calendar.d", "calendar.date"], "formula": "calendar.csv d -> date (one-to-one)", "unit": None, "assumptions": "No date outside calendar.csv"},
                             period_start={"source_columns": ["calendar.date"], "formula": "= date", "unit": None, "assumptions": "Daily grain"},
                             period_end={"source_columns": ["calendar.date"], "formula": "= date", "unit": None, "assumptions": "Daily grain"},
                             source_row_id={"source_columns": ["(record position)", "d_* header"], "formula": "'<1-based data record>:<d column>'", "unit": None, "assumptions": "Wide-to-long lineage keeps the source column"},
                             price={"source_columns": ["sell_prices.csv sell_price", "calendar.wm_yr_wk"], "formula": "left join on (store_id, item_id, wm_yr_wk of the day); the weekly average price repeats on each day of its week",
                                    "unit": "USD", "unit_status": "SOURCE_METADATA", "assumptions": "A missing weekly price stays NULL (guide: not sold that week), never 0"},
                             currency={"source_columns": ["M5 guide"], "formula": "'USD' where price is present", "unit": "USD", "unit_status": "SOURCE_METADATA", "assumptions": EXTERNAL_UNIT_EVIDENCE["m5"]["price"]["evidence"]}),
                     unit_status="SOURCE_METADATA")
    e.record_mapping("product_master", {"product_id": "item_id", "category": "cat_id"}, source, {}, "static", "dataset_product",
                     derived=derived([("product_id_namespace", ns["product"])], category_path={"source_columns": ["cat_id", "dept_id"], "formula": "'cat_id=<cat>|dept_id=<dept>'", "unit": None, "assumptions": "Verified item -> dept -> cat function"}))
    e.record_mapping("location_master", {"location_id": "store_id | state_id", "region": "state_id"}, source, {}, "static", "dataset_location",
                     derived=derived([("location_id_namespace", ns["store"] + " | " + ns["state"])], location_type={"source_columns": [], "formula": "'store' for store_id, 'state' for state_id", "unit": None, "assumptions": "States exist only as SNAP covariate scopes"}))
    e.record_mapping("calendar_event", {"event_name": "event_name_1 | event_name_2", "event_type": "event_type_1 | event_type_2", "date": "date"}, rel(e, paths["calendar"]), {}, "daily", "dataset_calendar",
                     derived=derived(source_row_id={"source_columns": ["(record position)"], "formula": "'<record>:event_name_<slot>'", "unit": None, "assumptions": "One row per named event slot"}))
    e.record_mapping("covariate_series", {"covariate_value": "snap_CA | snap_TX | snap_WI", "date": "date"}, rel(e, paths["calendar"]), {"covariate_value": "snap_<state>"}, "daily", "state", "binary_indicator",
                     derived([("location_id_namespace", ns["state"])], covariate_name={"source_columns": [], "formula": "'snap'", "unit": None, "assumptions": "Guide: SNAP purchase allowed on the date"}),
                     unit_status="SOURCE_METADATA")
    submission = read_all(paths["sample_submission"])
    forecast_columns = [c for c in submission.columns if c != "id"]
    e.checks["duplicate_release_investigation"] = m5_validation_release_evidence(paths["sales_train_evaluation"], paths["sales_train_validation"])
    e.checks["excluded_sources"] = {
        rel(e, paths["sales_train_validation"]): "strict prefix release of the evaluation file (see duplicate_release_investigation)",
        rel(e, paths["sample_submission"]): {"rows": len(submission), "forecast_columns": len(forecast_columns),
                                             "all_values_zero": bool((submission[forecast_columns].astype(float) == 0).all().all()),
                                             "reason": "Kaggle scoring template (F1..F28 placeholders), not an observation"},
        rel(e, paths["sell_prices"]) + " (beyond sales horizon)": {"records": price["records_beyond_horizon"],
                                                                  "reason": "weeks after d_1941 have no sales day to attach to; kept in raw, not converted"}}
    e.checks["price_availability"] = {
        "rows_with_weekly_price": price["priced_rows"], "zero_sales_rows_without_price": flags["weekly_price_absent"],
        "positive_sales_rows_without_price": flags["sales_without_weekly_price"],
        "price_range": [float(pd.to_numeric(prices["sell_price"]).min()), float(pd.to_numeric(prices["sell_price"]).max())],
        "interpretation": "Guide: a missing weekly price means the product was not sold that week. Every unpriced row has zero sales (0 contradictions); the price stays NULL (weekly_price_absent, informational), never 0."}
    e.checks["missing_vs_zero"] = {"source_cells": conservation["source_cells"], "null_cells": conservation["source_cells"] - conservation["source_nonnull"],
                                   "zero_cells": conservation["source_zero"], "canonical_zero_rows": conservation["canonical_zero"],
                                   "policy": "A zero cell is an observed zero sale; a NULL cell would stay NULL with missing:sales_qty (none exists in the release). No row exists for a date outside d_1..d_1941."}
    e.checks["hierarchy"] = {"items": int(first["item_id"].nunique()), "stores": int(first["store_id"].nunique()), "states": int(first["state_id"].nunique()),
                             "series": len(first), "functions_verified": ["item_id -> dept_id", "dept_id -> cat_id", "store_id -> state_id"]}
    e.checks["identity_namespaces"] = {"product": ns["product"], "location": [ns["store"], ns["state"]]}
    e.checks["unit_evidence"] = EXTERNAL_UNIT_EVIDENCE["m5"]
    e.sources += [{"file": source, "row_count": len(first), "canonical_rows": conservation["rows_canonical"], "status": "full file, wide d_1..d_1941 unpivoted", "family": "sales"},
                  {"file": rel(e, paths["sell_prices"]), "row_count": len(prices), "status": "joined to demand_series.price by (store, item, week)", "family": "price"},
                  {"file": rel(e, paths["calendar"]), "row_count": len(calendar), "canonical_rows": len(events) + len(snap), "status": "events and SNAP indicators", "family": "calendar"}]
    e.notes += ["Actual Walmart daily unit sales per store x item (10 stores, 3,049 items, 2011-01-29..2016-05-22). No on-hand inventory, cost, lead time, capacity or network exists; none is generated.",
                "sales_qty is observed unit sales (not demand): stock-outs are unobserved, so zero sales can be censored demand.",
                "Weekly sell_price is joined to every day of its week; a missing weekly price stays NULL (weekly_price_absent) and is never 0.",
                "Calendar events and state SNAP indicators live in calendar_event/covariate_series; M5 has no promotion field, so promotion stays NULL."]


# ---------------------------------------------------------------- Favorita

def run_favorita(e):
    prepare(e, "favorita")
    folder = e.folder / "raw/favorita_grocery_sales_forecasting/extracted/favorita-grocery-sales-forecasting"
    path = {n: folder / f"{n}.csv" for n in ("train", "test", "stores", "items", "transactions", "oil", "holidays_events", "sample_submission")}
    stores, items = read_all(path["stores"]), read_all(path["items"])
    if not (stores["store_nbr"].is_unique and items["item_nbr"].is_unique):
        raise ValueError("Favorita master identifiers repeat")
    if items.groupby("class")["family"].nunique().gt(1).any():
        raise ValueError("Favorita class does not nest in one family")
    locations, products = adapt_favorita_masters(stores, items, rel(e, path["stores"]), rel(e, path["items"]))
    e.write("location_master", locations, masters=False)
    e.write("product_master", products, masters=False)
    families = items.set_index("item_nbr")["family"]
    source = rel(e, path["train"])
    e.bulk_key_tables.add("demand_series")
    sales = StreamConservation(source, "demand_series", "sales_qty", "unit_sales", "fsum of every record, signed (returns included)")
    stats, promo_src, promo_can = Counter(), Counter(), Counter()
    negative_sum, dates, items_seen, stores_seen = [], set(), set(), set()
    fractional_items, negative_items = set(), set()
    null_promo_dates, previous_last, monotonic = [], None, True
    for chunk in arrow_csv_chunks(path["train"]):
        table, frame, measures = adapt_favorita_train(chunk, source, families)
        if len(frame) != len(chunk):
            raise AssertionError("Favorita record conservation failure")
        e.write(table, frame, masters=False)
        sales.add(chunk["unit_sales"], frame["sales_qty"], len(chunk), len(frame))
        qty = pd.to_numeric(chunk["unit_sales"], errors="coerce")
        ids = pd.to_numeric(chunk["id"]).to_numpy()
        stats.update(rows=len(chunk), ids_out_of_sequence=int((ids != np.arange(chunk.index[0], chunk.index[-1] + 1)).sum()),
                     negative_rows=int(qty.lt(0).sum()), zero_rows=int(qty.eq(0).sum()), fractional_rows=int((qty % 1).ne(0).sum()))
        negative_sum.append(exact_sum(qty[qty.lt(0)]))
        negative_items.update(chunk.loc[qty.lt(0), "item_nbr"].unique())
        fractional_items.update(chunk.loc[(qty % 1).ne(0), "item_nbr"].unique())
        promo_src.update(chunk["onpromotion"].fillna("<NULL>").value_counts().to_dict())
        promo_can.update(frame["promotion"].fillna("<NULL>").value_counts().to_dict())
        if chunk["onpromotion"].isna().any():
            null_promo_dates += [chunk.loc[chunk["onpromotion"].isna(), "date"].min(), chunk.loc[chunk["onpromotion"].isna(), "date"].max()]
        monotonic &= bool(chunk["date"].is_monotonic_increasing) and (previous_last is None or previous_last <= chunk["date"].iloc[0])
        previous_last = chunk["date"].iloc[-1]
        dates.update(chunk["date"].unique())
        items_seen.update(chunk["item_nbr"].unique())
        stores_seen.update(chunk["store_nbr"].unique())
    e.conservation.append(sales.result())
    promo_map = {"True": "true", "False": "false", "<NULL>": "<NULL>"}
    promo_ok = all(promo_can.get(promo_map[k], 0) == n for k, n in promo_src.items())
    e.conservation.append({"source_file": source, "table": "demand_series", "field": "promotion", "source_column": "onpromotion",
                           "source_counts": dict(promo_src), "canonical_counts": dict(promo_can), "passed": bool(promo_ok),
                           "method": "True->true, False->false, NULL->NULL; counts equal per value"})
    if not promo_ok:
        raise AssertionError("Favorita promotion conservation failure")
    context = {}
    for name, adapter, measure in [("transactions", adapt_favorita_transactions, "transactions"), ("oil", adapt_favorita_oil, "dcoilwtico")]:
        raw = read_all(path[name])
        frame = adapter(raw, rel(e, path[name]))
        e.write("covariate_series", frame, masters=False)
        check = StreamConservation(rel(e, path[name]), "covariate_series", "covariate_value", measure, "fsum and non-null count; NULL stays NULL")
        check.add(raw[measure], frame["covariate_value"], len(raw), len(frame))
        e.conservation.append(check.result())
        context[name] = raw
    holidays = read_all(path["holidays_events"])
    events = adapt_favorita_holidays(holidays, rel(e, path["holidays_events"]))
    e.write("calendar_event", events, masters=False)
    e.conservation.append({"source_file": rel(e, path["holidays_events"]), "table": "calendar_event", "field": "event_name", "source_column": "description",
                           "rows_source": len(holidays), "rows_canonical": len(events), "passed": len(holidays) == len(events)})
    test = Counter()
    test_dates, test_ids = set(), []
    for chunk in arrow_csv_chunks(path["test"]):
        test.update(rows=len(chunk))
        test.update(chunk["onpromotion"].fillna("<NULL>").value_counts().to_dict())
        test_dates.update(chunk["date"].unique())
        test_ids += [int(chunk["id"].iloc[0]), int(chunk["id"].iloc[-1])]
    submission = read_all(path["sample_submission"], chunksize=1_000_000)
    all_dates = pd.date_range(min(dates), max(dates)).strftime("%Y-%m-%d")
    per = items.set_index("item_nbr")["perishable"]
    ns = NAMESPACES["favorita"]
    e.checks["signed_sales_investigation"] = {
        "negative_rows": stats["negative_rows"], "negative_sum": math.fsum(negative_sum), "items_with_negative": len(negative_items),
        "items_with_negative_by_perishable": per.reindex(sorted(negative_items)).value_counts().to_dict(),
        "specification": "Kaggle data description: 'Negative values of unit_sales represent returns of that particular item.'",
        "policy": "Signed net unit_sales kept in sales_qty unchanged: negative:sales_qty (NEGATIVE_QUANTITY, blocking) + documented_return (informational). Never abs(), clipped or zeroed; returns_qty is not derived because the source records one signed value per item-store-day."}
    e.checks["zero_sales_absence"] = {
        "zero_rows": stats["zero_rows"], "calendar_days": len(all_dates), "days_with_rows": len(dates), "days_without_any_row": sorted(set(all_dates) - dates),
        "specification": "Kaggle data description: the training data does not include rows for items that had zero unit_sales for a store/date combination; stock availability is not recorded.",
        "policy": "An absent (date, store, item) is neither a zero sale nor a stock-out; no row is generated. The four absent days are 25 December (no store reported)."}
    e.checks["promotion_investigation"] = {
        "source_counts": dict(promo_src), "null_share": promo_src["<NULL>"] / stats["rows"],
        "null_date_range": [min(null_promo_dates), max(null_promo_dates)] if null_promo_dates else None,
        "specification": "Kaggle data description: 'Approximately 16% of the onpromotion values in this file are NaN.'",
        "policy": "NULL promotion stays NULL (promotion_not_reported, informational), never False."}
    e.checks["unit_investigation"] = {
        "fractional_rows": stats["fractional_rows"], "items_with_fractional_sales": len(fractional_items),
        "items_with_fractional_by_perishable": per.reindex(sorted(fractional_items)).value_counts().to_dict(),
        "conclusion": "Items sold by weight report fractional unit_sales; no field labels which items are weighed, so an integer-only item is not proof of a count unit. Quantity unit stays NULL (UNKNOWN); forecasting is valid per item series, never pooled across items."}
    e.checks["record_integrity"] = {
        "rows": stats["rows"], "ids_out_of_sequence": stats["ids_out_of_sequence"], "date_monotonic": monotonic, "date_range": [min(dates), max(dates)],
        "items_observed": len(items_seen), "items_not_in_items_csv": sorted(items_seen - set(items["item_nbr"])),
        "master_items_never_sold": len(set(items["item_nbr"]) - items_seen), "stores_observed": len(stores_seen),
        "stores_not_in_stores_csv": sorted(stores_seen - set(stores["store_nbr"]))}
    e.checks["excluded_sources"] = {
        rel(e, path["test"]): {"rows": test["rows"], "date_range": [min(test_dates), max(test_dates)], "id_range": [min(test_ids), max(test_ids)],
                               "onpromotion": {k: v for k, v in test.items() if k != "rows"},
                               "reason": "Kaggle scoring frame: unit_sales withheld, so no observation exists; its known-in-advance promotions stay in raw."},
        rel(e, path["sample_submission"]): {"rows": len(submission), "all_zero": bool(pd.to_numeric(submission["unit_sales"]).eq(0).all()), "reason": "scoring template"}}
    e.checks["identity_namespaces"] = {"product": ns["product"], "location": ns["store"]}
    e.checks["unit_evidence"] = EXTERNAL_UNIT_EVIDENCE["favorita"]
    e.record_mapping("demand_series", {"location_id": "store_nbr", "product_id": "item_nbr", "sales_qty": "unit_sales", "raw_date": "date"}, source, {"sales_qty": "unit_sales"},
                     "daily", "store_product", None,
                     derived([("product_id_namespace", ns["product"]), ("location_id_namespace", ns["store"])],
                             date={"source_columns": ["date"], "formula": "strict ISO date parse", "unit": None, "assumptions": "No date correction"},
                             category={"source_columns": ["items.family"], "formula": "item_nbr -> items.csv family", "unit": None, "assumptions": "Verified unique item_nbr"},
                             promotion={"source_columns": ["onpromotion"], "formula": "True->'true', False->'false', NULL->NULL", "unit": None, "assumptions": "Undocumented tokens fail loudly"}),
                     unit_status="UNKNOWN")
    e.record_mapping("location_master", {"location_id": "store_nbr", "city": "city", "region": "state", "location_subtype": "type", "location_cluster": "cluster"},
                     rel(e, path["stores"]), {}, "static", "dataset_location", derived=derived([("location_id_namespace", ns["store"])]))
    e.record_mapping("product_master", {"product_id": "item_nbr", "category": "family"}, rel(e, path["items"]), {}, "static", "dataset_product",
                     derived=derived([("product_id_namespace", ns["product"])],
                                     category_path={"source_columns": ["family", "class"], "formula": "'family=<f>|class=<c>'", "unit": None, "assumptions": "Verified class -> family function"},
                                     perishable={"source_columns": ["perishable"], "formula": "1->'true', 0->'false'", "unit": None, "assumptions": "Kaggle item metadata"}))
    e.record_mapping("covariate_series", {"location_id": "store_nbr", "covariate_value": "transactions | dcoilwtico"}, rel(e, path["transactions"]) + " | " + rel(e, path["oil"]),
                     {"covariate_value": "transactions | dcoilwtico"}, "daily", "store | national", "transaction (transactions) / NULL (oil)",
                     derived([("location_id_namespace", ns["store"])], covariate_name={"source_columns": [], "formula": "'transactions' | 'dcoilwtico'", "unit": None, "assumptions": "Source column name"}),
                     unit_status="SOURCE_METADATA (transactions) / UNKNOWN (oil)")
    e.record_mapping("calendar_event", {"event_name": "description", "event_type": "type", "event_locale": "locale", "event_locale_name": "locale_name", "raw_date": "date"},
                     rel(e, path["holidays_events"]), {}, "daily", "dataset_calendar",
                     derived=derived(event_transferred={"source_columns": ["transferred"], "formula": "True->'true', False->'false'", "unit": None,
                                                        "assumptions": "Kaggle: a transferred holiday officially falls on that day but was moved; the celebrated day is the matching Transfer row"}))
    e.sources += [{"file": source, "row_count": stats["rows"], "status": "full file streamed in 64 MB blocks", "family": "sales"},
                  *({"file": rel(e, path[n]), "row_count": len(context[n]), "status": "full file", "family": "covariate"} for n in ("transactions", "oil")),
                  {"file": rel(e, path["holidays_events"]), "row_count": len(holidays), "status": "full file", "family": "calendar"},
                  {"file": rel(e, path["stores"]), "row_count": len(stores), "status": "location master"},
                  {"file": rel(e, path["items"]), "row_count": len(items), "status": "product master"}]
    e.notes += ["Actual Corporación Favorita (Ecuador) daily store x item unit sales 2013-01-01..2017-08-15; no inventory, price, cost or network exists; none is generated.",
                "unit_sales is signed: negative values are documented returns (kept, NEGATIVE_QUANTITY + documented_return).",
                "Zero-sales rows are absent from the source; an absent key is not a zero and no zero rows are generated.",
                "Quantity unit is item-dependent (count or kg) and unlabelled: unit NULL, UNKNOWN_UNIT on every sales row.",
                "onpromotion NULL stays NULL (promotion_not_reported), never False.",
                "holidays_events locale_name is a city/state/country name, not a store; it is not joined to location_master."]


# ---------------------------------------------------------------- FreshRetailNet

def run_freshretailnet(e, batch_size=250_000):
    import pyarrow as pa
    import pyarrow.parquet as pq
    prepare(e, "freshretailnet")
    e.bulk_key_tables.add("demand_series")
    firsts, hierarchy, evidence = [], [], {}
    for name in ("train.parquet", "eval.parquet"):
        path = e.folder / "raw" / name
        source = rel(e, path)
        checks = {field: StreamConservation(source, "demand_series", field, column, "fsum and non-null count per record")
                  for field, column in [("sales_qty", "sale_amount"), ("stockout_hours", "stock_hour6_22_cnt"), ("discount_rate", "discount")]}
        tally, scalar_parts, hourly_sums, position = Counter(), [], [], 0
        for batch in pq.ParquetFile(path).iter_batches(batch_size=batch_size):
            table = pa.Table.from_batches([batch])
            frame, scalars, status, hourly = adapt_frn_sales(table, source, position)
            e.write("demand_series", frame, masters=False)
            for field, check in checks.items():
                check.add(scalars[check.entry["source_column"]], frame[field], len(scalars), len(frame))
            lo, hi = FRN_STOCKOUT_WINDOW
            tally.update(rows=len(frame), out_of_stock_hours_all_day=int(status.sum()), out_of_stock_hours_window=int(status[:, lo:hi].sum()),
                         sale_hours=int((hourly > 0).sum()), sale_hours_while_out_of_stock=int(((hourly > 0) & (status == 1)).sum()),
                         rows_all_window_hours_out_of_stock=int((status[:, lo:hi].sum(axis=1) == hi - lo).sum()),
                         rows_all_window_out_with_sales=int(((status[:, lo:hi].sum(axis=1) == hi - lo) & (scalars["sale_amount"].to_numpy() > 0)).sum()),
                         discount_zero=int(scalars["discount"].eq(0).sum()), discount_above_one=int(scalars["discount"].gt(1).sum()),
                         stockout_count_mismatch=int(frame["validation_flags"].str.contains("stockout_count_mismatch").sum()),
                         hourly_sales_sum_mismatch=int(frame["validation_flags"].str.contains("hourly_sales_sum_mismatch").sum()))
            hourly_sums.append((math.fsum(hourly[status == 1]), math.fsum(hourly.ravel())))
            scalar_parts.append(scalars[["store_id", "product_id", "dt", *FRN_WEATHER, "holiday_flag"]])
            ids = scalars.drop_duplicates("product_id").assign(_file=source, _record=lambda d: (d.index + 1).astype(str))
            firsts.append(ids)
            firsts.append(scalars.drop_duplicates("store_id").assign(_file=source, _record=lambda d: (d.index + 1).astype(str)))
            hierarchy.append(scalars[["product_id", *FRN_CATEGORY_LEVELS]].drop_duplicates())
            hierarchy.append(scalars[["store_id", "city_id"]].drop_duplicates())
            position += len(frame)
        for check in checks.values():
            e.conservation.append(check.result())
        scalars_all = pd.concat(scalar_parts)
        covariates = adapt_frn_covariates(scalars_all, source)
        e.write("covariate_series", covariates, masters=False)
        represented = covariates.groupby("covariate_name")["source_record_count"].sum()
        e.conservation.append({"source_file": source, "table": "covariate_series", "field": "source_record_count", "source_column": "precpt|avg_temperature|avg_humidity|avg_wind_level|holiday_flag",
                               "rows_source": len(scalars_all), "records_represented_per_covariate": {k: int(v) for k, v in represented.items()},
                               "method": "each covariate collapses identical member values (verified) and lists every member record id",
                               "passed": bool(represented.eq(len(scalars_all)).all())})
        if not represented.eq(len(scalars_all)).all():
            raise AssertionError("FreshRetailNet covariate lineage failure")
        dates = scalars_all["dt"]
        in_stockout, total = (math.fsum(a for a, _ in hourly_sums), math.fsum(b for _, b in hourly_sums))
        evidence[source] = {**dict(tally), "date_range": [dates.min(), dates.max()], "dates": int(dates.nunique()),
                            "stores": int(scalars_all["store_id"].nunique()), "series": int(scalars_all.groupby(["store_id", "product_id"]).ngroups),
                            "repeated_store_product_date": int(scalars_all.duplicated(["store_id", "product_id", "dt"]).sum()),
                            "hourly_sales_share_in_out_of_stock_hours": in_stockout / total if total else None}
        del scalars_all, scalar_parts
    products_h = pd.concat([h for h in hierarchy if "product_id" in h]).drop_duplicates()
    stores_h = pd.concat([h for h in hierarchy if "city_id" in h and "product_id" not in h]).drop_duplicates()
    if products_h["product_id"].duplicated().any() or stores_h["store_id"].duplicated().any():
        raise ValueError("FreshRetailNet product -> category or store -> city is not a function")
    first = pd.concat(firsts)
    products, stores = adapt_frn_masters(first)
    e.write("product_master", products, masters=False)
    e.write("location_master", stores, masters=False)
    train, test = evidence["05_FreshRetailNet/raw/train.parquet"], evidence["05_FreshRetailNet/raw/eval.parquet"]
    ns = NAMESPACES["freshretailnet"]
    e.checks["split_investigation"] = {"per_file": evidence, "date_overlap": not (train["date_range"][1] < test["date_range"][0]),
                                       "policy": "train (90 days) and eval (7 days) are disjoint dated observations of the same 50,000 series; both are converted and source_file keeps the split."}
    e.checks["stockout_semantics"] = {
        "specification": "Data card: stock_hour6_22_cnt 'The number of out-of-stock hours between 6:00 and 22:00'; hours_stock_status 'The hourly out-of-stock status'.",
        "verification": "sum(hours_stock_status[6:22]) == stock_hour6_22_cnt for every row (stockout_count_mismatch rows: %d), so 1 = out of stock." % (train["stockout_count_mismatch"] + test["stockout_count_mismatch"]),
        "sales_in_out_of_stock_hours": "Sales occur in hours marked out of stock (status is hourly; stock can run out within the hour); reported, not flagged.",
        "inventory": "The stock level is not released (the report says hourly stock levels were tracked in the WMS); only availability is. No inventory_snapshot is created and no stock quantity is inferred."}
    e.checks["discount_investigation"] = {
        "specification": "Data card: 'The discount rate (1.0 means no discount, 0.9 means 10% off)'.",
        "zero_discount_rows": train["discount_zero"] + test["discount_zero"], "above_one_rows": train["discount_above_one"] + test["discount_above_one"],
        "policy": "Values preserved; 0 (discount_rate_zero) and >1 (discount_rate_above_one) are informational SOURCE_ANOMALY flags on the covariate, the sales observation stays usable."}
    e.checks["provenance_and_realness"] = {
        "card": "https://huggingface.co/datasets/Dingdong-Inc/FreshRetailNet-50K (CC BY 4.0): '50,000 store-product 90-day time series of detailed hourly sales data from 898 stores in 18 major cities, encompassing 865 perishable SKUs'.",
        "technical_report": "arXiv 2505.16319 section 3.1: 'integrates multi-source operational data from 898 stores across 18 major Chinese cities ... (March to June 2024)'; abstract states 863 perishable SKUs.",
        "sku_count_discrepancy": "Card 865 vs abstract 863; the files hold %d product ids. Perishability is a dataset-level statement, so product_master.perishable stays NULL." % products_h["product_id"].nunique(),
        "classification": "REAL_SALES: actual operational sales with encoded ids and globally normalised quantities (not synthetic)."}
    e.checks["identity_namespaces"] = {"product": ns["product"], "location": ns["store"]}
    e.checks["unit_evidence"] = EXTERNAL_UNIT_EVIDENCE["freshretailnet"]
    source = "05_FreshRetailNet/raw/train.parquet | 05_FreshRetailNet/raw/eval.parquet"
    e.record_mapping("demand_series", {"location_id": "store_id", "product_id": "product_id", "category": "management_group_id", "sales_qty": "sale_amount", "discount_rate": "discount",
                                       "stockout_hours": "stock_hour6_22_cnt", "raw_date": "dt"},
                     source, {"sales_qty": "sale_amount", "stockout_hours": "stock_hour6_22_cnt", "discount_rate": "discount"}, "daily", "store_product", "normalized_sales_amount",
                     derived([("product_id_namespace", ns["product"]), ("location_id_namespace", ns["store"])],
                             promotion={"source_columns": ["activity_flag"], "formula": "1->'true', 0->'false'", "unit": None, "assumptions": "Card: 'Activity indicator'; report 3.1: marketing campaigns annotated. Not a price."},
                             stockout_window_hours={"source_columns": [], "formula": "22 - 6 = 16", "unit": "hour", "unit_status": "SOURCE_METADATA", "assumptions": "Card: out-of-stock hours between 6:00 and 22:00"},
                             hourly_stockout_status={"source_columns": ["hours_stock_status"], "formula": "24 chars '0'/'1', hour 0..23", "unit": None, "assumptions": "1 = out of stock (verified)"},
                             hourly_sales={"source_columns": ["hours_sale"], "formula": "24 values joined by '|', shortest round-trip text", "unit": "normalized_sales_amount", "unit_status": "SOURCE_METADATA", "assumptions": "Round-trip to identical doubles verified per batch"}),
                     unit_status="SOURCE_METADATA")
    e.record_mapping("covariate_series", {"location_id": "store_id", "covariate_value": "precpt | avg_temperature | avg_humidity | avg_wind_level | holiday_flag", "raw_date": "dt"},
                     source, {"covariate_value": "weather columns | holiday_flag"}, "daily", "store (weather) | national (holiday_flag)", "NULL (weather) / binary_indicator (holiday)",
                     derived([("location_id_namespace", ns["store"])],
                             source_row_id={"source_columns": ["(record position)"], "formula": "every member record id joined by '|'", "unit": None, "assumptions": "Members carry identical values (verified, else fail)"},
                             source_record_count={"source_columns": ["(record position)"], "formula": "member record count", "unit": None, "assumptions": "Collapse of identical repeated attributes, not a sum"}),
                     unit_status="UNKNOWN (weather) / SOURCE_METADATA (holiday)")
    e.record_mapping("product_master", {"product_id": "product_id", "category": "management_group_id"}, source, {}, "static", "dataset_product",
                     derived=derived([("product_id_namespace", ns["product"])], category_path={"source_columns": list(FRN_CATEGORY_LEVELS), "formula": "'level=value|...'", "unit": None, "assumptions": "Verified product -> category function"}))
    e.record_mapping("location_master", {"location_id": "store_id", "city": "city_id"}, source, {}, "static", "dataset_location", derived=derived([("location_id_namespace", ns["store"])]))
    e.sources += [{"file": f, "row_count": v["rows"], "date_range": v["date_range"], "status": "full parquet, 250k-row batches"} for f, v in evidence.items()]
    e.notes += ["Actual Dingdong fresh-retail sales (898 stores, 18 Chinese cities, 865 product ids, 2024-03-28..2024-07-02) with encoded ids.",
                "sale_amount is globally normalised by an undisclosed coefficient: unit normalized_sales_amount, not items or currency.",
                "hours_stock_status 1 = out of stock (verified against stock_hour6_22_cnt). Availability is not inventory: no inventory_snapshot.",
                "Weather is store-day and holiday_flag is date-level; both are collapsed from identical member rows with every record id kept."]


# ---------------------------------------------------------------- KAMP

KAMP_GUIDEBOOK_QUOTES = {
    "공사": "발주공사명", "부위": "발주공사부위", "출하일자": "자재출하일자", "grade_columns": "해당강종출하량",
    "collection_period": "2022년 10월 01일 ~ 2023년 9월 30일", "file": "I_read.xlsx", "size": "9,893개(변수 13개 x Raw 761개)",
    "stock_source_not_released": "금일 실물 재고 수집 수량 데이터", "shipment_basis": "발주DATA(바리스트)",
}


def read_kamp_workbook(path):
    import openpyxl
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = book.sheetnames
        sheet = book[sheets[0]]
        rows = list(sheet.iter_rows(min_row=1, values_only=True))
    finally:
        book.close()
    header = [str(v) if v is not None else None for v in rows[0]]
    data = [(number, values) for number, values in enumerate(rows[1:], start=2) if any(v is not None for v in values)]
    raw = pd.DataFrame([v for _, v in data], columns=header)
    return sheets, raw, [n for n, _ in data]


def run_kamp(e):
    prepare(e, "kamp")
    path = next((e.folder / "raw").glob("*.xlsx"))
    source = rel(e, path)
    sheets, raw, sheet_rows = read_kamp_workbook(path)
    frame, long = adapt_kamp_shipments(raw, source, sheets[0], sheet_rows)
    e.write("inventory_flow", frame)
    by_column = frame["source_row_id"].str.split(":").str[1]
    for column in [*KAMP_GRADES, KAMP_TOTAL]:
        check = StreamConservation(source, "inventory_flow", "shipment_qty", column, "fsum of the source column vs canonical rows with that column id")
        check.add(raw[column], frame.loc[by_column.eq(column), "shipment_qty"], len(raw), int(by_column.eq(column).sum()))
        e.conservation.append(check.result())
    grades = raw[list(KAMP_GRADES)].apply(pd.to_numeric)
    identity = grades.sum(axis=1) - pd.to_numeric(raw[KAMP_TOTAL])
    e.conservation.append({"source_file": source, "table": "inventory_flow", "field": "(all cells)", "source_column": "10 quantity columns",
                           "rows_source": len(raw), "source_cells": len(raw) * (len(KAMP_GRADES) + 1), "rows_canonical": len(frame),
                           "passed": len(frame) == len(raw) * (len(KAMP_GRADES) + 1)})
    guide = (e.folder / "processed/guidebook_text.txt").read_text(encoding="utf-8")
    flat = " ".join(guide.split())
    e.checks["guidebook_semantics"] = {k: {"quote": v, "found": v in guide or v in flat} for k, v in KAMP_GUIDEBOOK_QUOTES.items()}
    e.checks["guidebook_semantics"]["unit"] = "The variable table's 단위 column is empty for every variable; no unit is printed in the workbook header."
    e.checks["row_total_identity"] = {"rows": len(raw), "rows_equal_within_1e-9": int(identity.abs().lt(1e-9).sum()), "max_abs_difference": float(identity.abs().max()),
                                      "policy": "합계 is kept as a separately scoped all-products row (aggregate_product_total, OUT_OF_SCOPE); never summed with grade rows."}
    e.checks["zero_vs_null"] = {"grade_cells": int(grades.size), "zero_cells": int(grades.eq(0).sum().sum()), "null_cells": int(raw[list(KAMP_GRADES)].isna().sum().sum()),
                                "zero_by_grade": grades.eq(0).sum().astype(int).to_dict(), "all_zero_rows": int(grades.sum(axis=1).eq(0).sum()),
                                "policy": "A zero cell is an observed zero shipment of that grade in that record (guidebook: Not null); a NULL cell would stay NULL with missing:shipment_qty."}
    dates = pd.to_datetime(raw["출하일자"].astype(str), format="%Y/%m/%d", errors="coerce")
    e.checks["project_part_semantics"] = {
        "projects": raw["공사"].value_counts().to_dict(), "parts": int(raw["부위"].nunique()),
        "repeated_project_part_date": int(raw.duplicated(list(KAMP_KEYS)).sum()), "fully_identical_rows": int(raw.duplicated().sum()),
        "date_range": [str(dates.min().date()), str(dates.max().date())], "shipment_dates": int(dates.nunique()), "invalid_dates": int(dates.isna().sum()),
        "conclusion": "공사 is the ordering construction project (masked name) and 부위 a building member/zone of that project: project_id/project_part, not locations. The shipping plant is not keyed, so location_id stays NULL (UNKNOWN_LOCATION); no location_master is created."}
    e.checks["no_inventory"] = {"workbook_columns": list(raw.columns),
                                "conclusion": "The guidebook lists a same-day physical stock collection (MS-SQL), but the released workbook has shipment columns only. No inventory_snapshot is created and shortage/surplus is not derived."}
    e.checks["identity_namespaces"] = {"product": NAMESPACES["kamp"]["product"]}
    e.checks["unit_evidence"] = EXTERNAL_UNIT_EVIDENCE["kamp"]
    e.record_mapping("inventory_flow", {"project_id": "공사", "project_part": "부위", "raw_date": "출하일자", "shipment_qty": "HD10..UHD25 | 합계", "product_id": "grade column header", "product_name": "grade column header"},
                     source, {"shipment_qty": "HD10..UHD25 | 합계"}, "event", "shipper_outbound_by_order_project", None,
                     derived([("product_id_namespace", NAMESPACES["kamp"]["product"])],
                             date={"source_columns": ["출하일자"], "formula": "strict '%Y/%m/%d' parse", "unit": None, "assumptions": "Event date; missing days are not zero-filled"},
                             source_row_id={"source_columns": ["(worksheet row)", "column header"], "formula": "'<worksheet row>:<column>'", "unit": None, "assumptions": "Wide-to-long lineage keeps the source column"},
                             product_grain={"source_columns": ["column header"], "formula": "'all_products' for 합계, else 'product'", "unit": None, "assumptions": "Reported row total, separately scoped"},
                             flow_kind={"source_columns": [], "formula": "'outbound_shipment_to_order_project'", "unit": None, "assumptions": "Guidebook: 발주(바리스트) 기반 출하량"}),
                     unit_status="UNKNOWN")
    e.sources.append({"file": source, "sheet_or_table": sheets[0], "sheets": sheets, "row_count": len(raw), "canonical_rows": len(frame), "columns": list(raw.columns),
                      "worksheet_rows": [min(sheet_rows), max(sheet_rows)], "status": "full workbook"})
    e.notes += ["Actual 인스틸(주) MES order-based rebar shipments for one masked construction project (2022-10-11..2023-09-26), 9 grades.",
                "Wide grade columns unpivoted to one inventory_flow row per (worksheet row, grade) with '<row>:<grade>' lineage; 합계 kept as an all-products row.",
                "Quantity unit not stated (UNKNOWN). No stock, cost, plant location or network: no inventory_snapshot, location_master or transfer_network."]


# ---------------------------------------------------------------- validation matrix

FACT_TABLES = ("inventory_snapshot", "inventory_flow", "transfer_network", "demand_series", "location_product_observation")
LEVEL_RANK = {"UNSUPPORTED": 0, "BENCHMARK_ONLY": 1, "PARTIAL": 2, "FULL": 3}
INVENTORY_ALGORITHMS = ("Turnover", "Disposal Risk", "Safety Stock", "EOQ")
NETWORK_ALGORITHMS = ("Transport Cost", "Greedy", "VHS", "Pareto", "MILP", "Optimality Gap", "Varo Final")
MATRIX_NOTES = {
    "suhyup": "Daily center stock/flow (count unit UNKNOWN, kg DIRECT); routes/costs/constraints are the existing 31-day benchmark proxy, not observed transfers.",
    "logisall": "Masked ZIP zones, not facilities; national monthly stock cannot be joined to zone-daily sales; observed zone-to-zone distribution has no cost.",
    "jangbogo": "Warehouse x category sales and purchase requests; purchase requests are not receipts or demand; 4 negative requests kept.",
    "nfqs": "Quarterly regional cold-storage stock in tons (cooperating firms only); no flow, sales or network.",
    "aihub": "One unidentified site aggregate (61 days) plus item measurements; no location key.",
    "m5": "Actual daily unit sales + weekly prices; no inventory; validation release is a strict prefix (converted once). Demand Forecast FULL by recorded review.",
    "favorita": "Actual daily unit sales, signed (returns), zero-sales rows absent, unit item-dependent (UNKNOWN); promotions with 17% NULL kept NULL.",
    "freshretailnet": "Actual normalised sales with hourly stockout status (availability, not stock); units are a normalised amount.",
    "kamp": "Order-based rebar shipments for one project; 공사/부위 are project attributes, not locations; unit UNKNOWN; no stock.",
}


def _present(report, table, *columns):
    coverage = report.get("field_coverage", {}).get(table, {})
    return any(coverage.get(c, "MISSING") != "MISSING" for c in columns)


def _distinct(path, column):
    import pyarrow.parquet as pq
    if not path.exists():
        return 0
    return len(pq.read_table(path, columns=[column]).column(column).drop_null().unique())


def matrix_row(root, dataset, catalog):
    import pyarrow.parquet as pq
    folder = root / DATASETS[dataset]
    report_path, coverage_path = folder / "results/data_quality_report.json", folder / "results/algorithm_coverage.csv"
    entry = catalog.get(DATASETS[dataset].replace("\\", "/"), {})
    row = {"dataset": dataset, "catalog_id": entry.get("dataset_id"), "real_or_synthetic": entry.get("real_or_synthetic")}
    if not report_path.exists():
        return {**row, "notes": "not converted to canonical"}
    report = json.loads(report_path.read_text(encoding="utf-8"))
    coverage = pd.read_csv(coverage_path, encoding="utf-8-sig").set_index("algorithm")["support_level"]
    tables = sorted(p.stem.removeprefix("canonical_") for p in (folder / "processed").glob("canonical_*.parquet"))
    counts = {t: pq.ParquetFile(folder / f"processed/canonical_{t}.parquet").metadata.num_rows for t in tables}
    ranges = [v for t, v in report.get("date_ranges", {}).items() if t in FACT_TABLES]
    best = lambda names: max((coverage.get(n, "UNSUPPORTED") for n in names), key=LEVEL_RANK.__getitem__)
    level_lists = {level: "|".join(a for a in ALGORITHMS if coverage.get(a) == level) for level in ("FULL", "PARTIAL", "BENCHMARK_ONLY")}
    return {**row,
            "date_range": f"{min(r[0] for r in ranges)}..{max(r[1] for r in ranges)}" if ranges else None,
            "rows": sum(counts.get(t, 0) for t in FACT_TABLES),
            "products": _distinct(folder / "processed/canonical_product_master.parquet", "product_id"),
            "locations": _distinct(folder / "processed/canonical_location_master.parquet", "location_id"),
            "has_inventory": _present(report, "inventory_snapshot", "inventory_qty"),
            "has_sales": _present(report, "demand_series", "sales_qty") or _present(report, "inventory_flow", "sales_qty"),
            "has_demand": _present(report, "demand_series", "demand_qty") or _present(report, "inventory_flow", "demand_qty"),
            "has_price": _present(report, "demand_series", "price") or _present(report, "product_master", "sell_price", "unit_cost"),
            "has_promotion": _present(report, "demand_series", "promotion", "discount_rate"),
            "has_flow": _present(report, "inventory_flow", "inbound_qty", "outbound_qty", "shipment_qty", "order_qty") or _present(report, "transfer_network", "shipment_qty"),
            "has_source_target": _present(report, "transfer_network", "source_id") and _present(report, "transfer_network", "target_id"),
            "has_cost": _present(report, "transfer_network", "transport_cost") or _present(report, "product_master", "unit_cost", "holding_cost", "ordering_cost", "disposal_cost"),
            "has_capacity": any(_present(report, t, "capacity") for t in ("inventory_snapshot", "location_master", "transfer_network")),
            "forecast_ready": coverage.get("Demand Forecast", "UNSUPPORTED"), "inventory_ready": best(INVENTORY_ALGORITHMS),
            "network_ready": best(NETWORK_ALGORITHMS), "dqn_ready": coverage.get("DQN", "UNSUPPORTED"),
            "notes": MATRIX_NOTES.get(dataset, ""),
            "canonical_tables": "|".join(tables), "context_rows": sum(counts.get(t, 0) for t in ("calendar_event", "covariate_series")),
            "quantity_unit_status": "|".join(sorted({s for t in FACT_TABLES for s in report.get("semantic_checks", {}).get("unit_consistency", {}).get(t, {}).get("unit_status", {})})),
            "transform_version": report.get("transform_version"),
            "full_algorithms": level_lists["FULL"], "partial_algorithms": level_lists["PARTIAL"], "benchmark_only_algorithms": level_lists["BENCHMARK_ONLY"],
            "all_quantity_checks_passed": report.get("all_quantity_checks_passed")}


def write_validation_matrix(root=DATA_ROOT):
    """One row per canonical dataset, derived only from its canonical outputs, coverage and the collection catalog."""
    root = Path(root)
    catalog_path = root / "_COLLECTION_STATUS/MASTER_DATASET_CATALOG.csv"
    catalog = {}
    if catalog_path.exists():
        frame = pd.read_csv(catalog_path, encoding="utf-8-sig", dtype=str)
        catalog = {str(r["local_folder"]).replace("\\", "/"): r.to_dict() for _, r in frame.iterrows() if pd.notna(r["local_folder"])}
    matrix = pd.DataFrame([matrix_row(root, dataset, catalog) for dataset in DATASETS])
    out = root / "_COLLECTION_STATUS/EXTERNAL_VALIDATION_MATRIX.csv"
    out.parent.mkdir(parents=True, exist_ok=True)
    matrix.to_csv(out, index=False, encoding="utf-8-sig")
    return out


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--data-root", type=Path, default=DATA_ROOT)
    parser.add_argument("--matrix-only", action="store_true", help="rebuild EXTERNAL_VALIDATION_MATRIX.csv from existing outputs")
    args = parser.parse_args()
    if not args.matrix_only:
        parser.error("use `python -m services.canonical_data_pipeline --datasets ...` to convert; this entry point only rebuilds the matrix")
    print(write_validation_matrix(args.data_root))


if __name__ == "__main__":
    main()
